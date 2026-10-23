from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'architecture_tlm/memory_savings_20261023'
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.55);sec.bottom_margin=Inches(.55)
sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
styles=doc.styles
for name,size in [('Normal',11),('Title',22),('Heading 1',16),('Heading 2',12),('Caption',9)]:
    st=styles[name];st.font.name='Microsoft YaHei';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    st.paragraph_format.space_after=Pt(5)
    st.paragraph_format.line_spacing=1.05
    if name.startswith('Heading'):st.paragraph_format.space_before=Pt(6)
styles['Normal'].paragraph_format.widow_control=True
for st in styles:
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='概率计算架构的SRAM与DRAM节省总结'
doc.core_properties.subject='存储容量 读写流量 条件能量与证据边界'
doc.core_properties.author='jieruiz'
header=sec.header.paragraphs[0];header.text='pbit LLM　存储架构分析';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('2026年10月23日　·　');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def inline(p,text):
    parts=re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\(https?://[^)]+\))',text)
    for s in parts:
        if not s:continue
        m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',s)
        if m:
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'215D82');pr.append(c);r.append(pr)
            t=OxmlElement('w:t');t.text=m[1];r.append(t);h.append(r);p._p.append(h)
        else:
            r=p.add_run(s[2:-2] if s.startswith('**') else s);r.bold=s.startswith('**')

def table(lines):
    data=[[s.strip() for s in l.strip().strip('|').split('|')] for l in lines]
    data=[r for r in data if not all(re.fullmatch(r':?-+:?',c) for c in r)]
    n=len(data[0]);t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    first=data[0][0]
    if '改进'==first:widths=[1.1,2.05,1.9,2.05]
    elif first=='上下文与方法':widths=[3.25,1.25,1.3,1.3] if n==4 else [2.5,.85,.85,1.1,1.8]
    elif first=='优先级':widths=[.65,2.5,3.95]
    elif first=='模型':widths=[1.8,1.5,1.5,2.3]
    elif n==3:widths=[2.8,2.15,2.15]
    elif first.startswith('PR127'):widths=[2.3,1.0,1.2,2.6]
    else:widths=[2.8,1.4,1.5,1.4]
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,rowdata in enumerate(data):
        row=t.rows[0] if i==0 else t.add_row()
        if i==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for j,(cell,txt,w) in enumerate(zip(row.cells,rowdata,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE8EF' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF'));pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for pos,v in [('top',70),('bottom',70),('left',90),('right',90)]:
                el=OxmlElement('w:'+pos);el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.12
            inline(p,txt)
            if (j>0 and re.match(r'^[0-9]|^0 ',txt)) or j==0 and first=='优先级':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.font.size=Pt(9.5);r.bold=i==0 or r.bold
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1;p.add_run().font.size=Pt(2)

lines=(OUT/'REPORT_ZH.md').read_text(encoding='utf-8').splitlines();i=0;break_next=False
while i<len(lines):
    s=lines[i]
    if s.startswith('|'):
        ls=[]
        while i<len(lines) and lines[i].startswith('|'):ls.append(lines[i]);i+=1
        table(ls);continue
    if s=='<!-- pagebreak -->':break_next=True
    elif s.startswith('# '):inline(doc.add_paragraph(style='Title'),s[2:])
    elif s.startswith('## '):
        p=doc.add_paragraph(style='Heading 1');p.paragraph_format.page_break_before=break_next;break_next=False;inline(p,s[3:])
    elif s.startswith('### '):inline(doc.add_paragraph(style='Heading 2'),s[4:])
    elif s:
        p=doc.add_paragraph();inline(p,s[2:] if s.startswith('- ') else s)
        if s.startswith('- '):p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.08
    i+=1
target=OUT/'概率计算架构的SRAM与DRAM节省总结.docx';doc.save(target)
print(target)
