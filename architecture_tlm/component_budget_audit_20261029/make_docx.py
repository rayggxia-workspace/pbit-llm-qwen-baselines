"""Create a readable Word report from Issue210's reconciled JSON."""
from pathlib import Path
import argparse,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ap=argparse.ArgumentParser();ap.add_argument('--input',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);args=ap.parse_args()
d=json.loads(args.input.read_text(encoding='utf-8'));tot=d['totals']
doc=Document();sec=doc.sections[0]
sec.page_width,sec.page_height=Inches(8.5),Inches(11)
sec.left_margin=sec.right_margin=Inches(.65)
sec.top_margin=sec.bottom_margin=Inches(.62)
sec.header_distance=sec.footer_distance=Inches(.24)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
 st.paragraph_format.line_spacing=1.08;st.paragraph_format.space_after=Pt(5)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].paragraph_format.space_before=Pt(7)
doc.styles['Heading 2'].font.size=Pt(11);doc.styles['Heading 2'].font.bold=True
def p(txt='',bold=False,size=None,after=5):
 x=doc.add_paragraph();x.paragraph_format.space_after=Pt(after);r=x.add_run(txt);r.bold=bold
 if size:r.font.size=Pt(size)
 return x
def h(txt):doc.add_heading(txt,level=1)
def page(txt):
 x=doc.add_heading(txt,level=1);x.paragraph_format.page_break_before=True
def table(headers,rows,widths,fs=8.5):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 border=OxmlElement('w:tblBorders')
 for k in ['top','bottom','left','right','insideH','insideV']:
  e=OxmlElement('w:'+k)
  for a,v in [('val','single'),('sz','5'),('color','D9D9D9')]:e.set(qn('w:'+a),v)
  border.append(e)
 t._tbl.tblPr.append(border)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for i,vals in enumerate([headers]+rows):
  row=t.rows[0] if i==0 else t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for j,(cell,txt,w) in enumerate(zip(row.cells,vals,widths)):
   cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tc=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'29455C' if i==0 else ('F2F5F7' if i%2==0 else 'FFFFFF'));tc.append(sh)
   mar=OxmlElement('w:tcMar')
   for tag,num in [('top','52'),('bottom','52'),('left','65'),('right','65')]:
    e=OxmlElement('w:'+tag);e.set(qn('w:w'),num);e.set(qn('w:type'),'dxa');mar.append(e)
   tc.append(mar);x=cell.paragraphs[0];x.paragraph_format.space_after=Pt(0);x.paragraph_format.line_spacing=1.02
   x.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
   r=x.add_run(str(txt));r.font.size=Pt(fs);r.font.name='Microsoft YaHei';r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
   if i==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
 return t
def math(txt):
 x=p();x.alignment=WD_ALIGN_PARAGRAPH.CENTER
 e=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=txt;r.append(t);e.append(r);x._p.append(e)
def link(label,url):
 x=p();x.paragraph_format.space_after=Pt(4);hyper=OxmlElement('w:hyperlink');hyper.set(qn('r:id'),x.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'24597A');pr.append(c);r.append(pr);tx=OxmlElement('w:t');tx.text=label;r.append(tx);hyper.append(r);x._p.append(hyper)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
f.add_run('2026年10月29日  ·  Issue210  ·  已有TLM预算派生  ·  第 ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field);f.add_run(' 页').font.size=Pt(8)

doc.add_paragraph('整机TLM功耗面积与时间分项',style='Title')
p('用原预算定位各模块的优化空间',bold=True,size=11)
p('片上能量主要来自共享数字MAC和FFN计数器；面积主要在计数器、共享SRAM及MAC。FFN随机源与比较器在原账本里多为合并计价，数制转换也有尚未独立报价的部分。本报告按原记录拆分，保留这些边界。')
table(['原整机锚点','三seed平均结果'],[
 ['模型与任务','Qwen3-4B后训练形状；36层；B1；2K；1个decode token'],
 ['资源与算法','500MHz目标；32×64×64 SC；2048 MAC/拍；GU64/NL128'],
 ['Attention','摘要筛选＋准确Softmax＋PV系统采样；未启用Ising'],
 ['时间',f"{tot['TPOT_ms']:.6f} ms/token"],
 ['已价片上能量与功率',f"{tot['onchip_mJ']:.4f} mJ/token；{tot['onchip_W_common_mean_TPOT']:.6f} W"],
 ['已价面积',f"原值{tot['raw_priced_area_mm2']:.4f} mm²；×0.81后{tot['equivalent_55nm_priced_area_mm2']:.4f} mm²"],
 ],[2.0,5.2],fs=9)
h('总括占比')
table(['片上类别','mJ/token','平均W','能量占比','55nm面积\nmm²','面积占比'],[
 [x['group'],f"{x['energy_mJ_per_token']:.3f}",f"{x['average_power_W']:.6f}",f"{x['onchip_energy_share_pct']:.2f}%",f"{x['equivalent_55nm_area_mm2']:.4f}",f"{x['priced_area_share_pct']:.2f}%"] for x in d['coarse_groups']],
 [2.15,1.0,.95,1.0,1.1,1.0],fs=8.7)
p('占比的分母是已计价片上小计；同一TPOT下，平均功率占比等于能量占比。面积按共享硬件计一次，不按36层复制。未定价部分不因此成为零成本。',size=9.2)
p(f"外存另有旧粗预算{tot['external_proxy_W']:.3f}W，加总约{tot['conditional_system_proxy_W']:.3f}W。它使用64pJ/B＋50mW背景，约占该条件总能量71.0%，不是有效HBM功率标定。",bold=True,size=9.5)

page('二 原计价账本的全部模块')
p('以下14行能量与已价面积闭合到第一页小计。raw面积保留原值，55nm等效只乘0.81一次。最后一行没有独立面积价，不写为实际零面积。',size=9.5)
short={
 'FFN_counter':'FFN权重门控／符号／对齐／计数',
 'FFN_postscale_restore_proxy':'FFN公共系数与数值恢复',
 'ordinary_BF16_MACs':'共享数字MAC阵列',
 'PV_vector_adders':'PV向量累加器',
 'wide_partial_activity_proxy':'FFN跨tile宽部分和',
 'shared_SRAM':'共享SRAM 8MiB',
 'local_CIM_SRAM_base':'FFN驻留权重SRAM',
 'digital_local_buffers':'数字本地operand与激活缓冲',
 'DMA_response_buffers':'DMA响应缓冲',
 'FFN_shared64_input_frontend':'FFN入口随机源＋比较器',
 'FFN_nonlinear_10_local_PRNG_endpoints':'FFN非线性随机源＋比较器',
 'Attention_existing_PRNG_endpoints':'Attention随机端点',
 'FFN_and_selector_sigmoid_tables':'FFN和筛选器sigmoid表',
 'other_vector_and_nonlinear_proxy':'其他向量和非线性代理',
}
table(['物理计价模块','mJ/token','平均W','片上占比','raw面积\nmm²','55nm面积\nmm²'],[
 [short[x['module']],f"{x['energy_mJ_per_token']:.6f}",f"{x['average_power_W']:.8f}",f"{x['onchip_energy_share_pct']:.4f}%",'未独立价' if x['raw_priced_area_mm2'] is None else f"{x['raw_priced_area_mm2']:.5f}",'未独立价' if x['equivalent_55nm_area_mm2'] is None else f"{x['equivalent_55nm_area_mm2']:.5f}"] for x in d['hardware_modules']],
 [2.15,1.1,1.1,.95,.95,.95],fs=8.3)
h('怎样理解这张表')
p('计数器一行包括权重门控、符号、档位对齐、归约和时间累计；不能再给这些操作加一套相同单价。恢复核为旧组合核代理，跨tile宽累加另列。')
p('输入前端已包括LFSR、64路比较和分发；非线性10个Bivium端点已包括比较器。没有匹配的内部分项数据，无法把每个合并价格精确拆成“纯随机源”和“纯比较器”。')
p('普通MAC包括QKV、O、候选QK、摘要打分及lm_head。PV加法单独收费，其他向量功能只保留旧动作代理；RMSNorm、RoPE、exp和倒数没有完整同工艺阵列PPA。')
p('SRAM四行是不同物理存储，允许分别收费。共享SRAM不能重新按FFN、Attention各加一遍；同一权重经过DMA响应缓冲和本地阵列也不能只计其中一次。')

page('三 FFN随机计算的具体边界')
service={x['service']:x['service_busy_ms'] for x in d['overlapping_service_busy_times']}
table(['步骤','已有费用或缺项','记录忙时ms'],[
 ['x固定四档编码','独立PPA缺失；92,160值',f"{service['FFN入口编码']:.6f}"],
 ['共享随机门限与64比较器','0.007618mJ；已含入口随机源和分发','与GU采样重叠'],
 ['Wg/Wu采样及计数','合并计数核分担；64次采样',f"{service['FFN GU采样与计数']:.6f}"],
 ['g/u编码与sigmoid等','部分LUT已计；编码／符号指数独立PPA缺',f"{service['FFN非线性编码与查表']:.6f}"],
 ['非线性随机源及比较','10端点约1.549196mJ；200路覆盖192路','与Wd采样重叠'],
 ['三路AND直接送Wd','AND控制未独立闭合；Wd计数已计；128次采样',f"{service['FFN AND直连Wd采样与计数']:.6f}"],
 ['权重与系数编程','SRAM按装载与保持收费；并行外围未价',f"{service['FFN权重与系数编程']:.6f}"],
 ['旧归约与系数恢复宏','恢复代理约1.812209mJ；避免重复七桶费用',f"{service['FFN旧恢复宏']:.6f}"],
 ['最终宽定点转BF16','92,160值；独立转换功耗面积未价','0.011664'],
 ],[1.9,4.1,1.2],fs=8.6)
p('忙时是服务预约量，GU随机门限与计数、非线性随机源与Wd使用同一采样时段；编程也可能与计算重叠，不能把上表逐项相加成FFN总时间。FFN全阶段实际约25.210ms，包含供数等待。',size=9.5)
h('权重存储和中间结果')
p('原FFN权重固定W8，两个本地权重bank合计256KiB；只按权重编程写入与保持计价，没有按64／128拍反复收普通SRAM完整读取费。定制CIM全并行读出、节点负载与广播驱动仍未计价。新双随机W侧SNG不属于本锚点。')
p('g/u恢复后再进行P8编码，sigmoid由真实g的signed8驱动；三路AND形成比特流直接进入Wd，最终输出才宽累加、舍入回BF16。不能用sigmoid表这一行的小能量代表完整随机非线性。')
h('时间模型和能量模型有一处边界尚未统一')
p('能量后处理已采用合并计数器并排除重复七桶归约；原时序仍保留旧SC_bucket_shift_reduce和系数恢复宏。因此恢复忙时不能拿来追加七套计数电路面积。本报告保持旧值，不把尚未执行的集成优化算成收益。')
h('数制转换的缺项')
p('数字矩阵输出另有313,344次FP32→BF16，忙时0.059904ms。它与FFN最终92,160次宽定点→BF16合计约0.071568ms，已在时间表里，却没有独立能耗面积报价；x及g/u编码也有类似缺口。')

page('四 Attention各环节与共享数字MAC')
p('QKV和O投影仍用普通数字MAC。当前筛选是摘要点积、max4归一化、sigmoid概率保留与固定预算修整，不是耦合p-bit网络退火；之后准确Softmax，再用PV系统采样和GQA地址并集读V。')
table(['功能','MAC次数','有效计算增量mJ'],[
 [x['function'],f"{x['MACs']:,.0f}",f"{x['work_increment_mJ']:.6f}"] for x in d['functional_MAC_work']],
 [2.7,2.4,2.1],fs=9)
p(f"这五项合计1,427,046,400次MAC，只解释约7.035814mJ的有效动作增量。共享MAC账目余下{d['shared_MAC_background_mJ']:.6f}mJ是背景动态和漏电，其中背景动态也包含工作期间基底。MAC利用率约{d['shared_MAC_utilization_pct']:.3f}%。不能把全部69.532mJ都归为Attention投影计算。",size=9.5)
sel=['Attention QKV投影算术','Attention O投影算术','RoPE算术','摘要筛选算术','K读入与片存搬运','QK打分与缩放','准确Softmax含输入输出SRAM','PV采样选址与GQA合并','V外存读取 片存装载与广播','V选择累加与平均','Attention随机服务真实响应']
table(['环节','服务忙时ms','费用归属'],[
 [x,f"{service[x]:.6f}",('共享MAC' if '投影' in x or x=='QK打分与缩放' else '旧向量代理／共享资源' if x in ['RoPE算术','摘要筛选算术','准确Softmax含输入输出SRAM'] else '外存粗预算＋SRAM' if x in ['K读入与片存搬运','V外存读取 片存装载与广播'] else 'PV累加器' if '累加' in x else '随机端点及控制／共享存储')] for x in sel],
 [3.7,1.2,2.3],fs=8.5)
p('服务忙时跨行可重叠；准确Softmax的exp／倒数仍用旧向量与非线性单价，部分筛选／选址控制没有独立PPA。不能从旧代理成本很小推断实电路同样小。',size=9.5)
p('Ising Softmax：本锚点未启用，因此没有可从2.516W里拆出的Ising占比。新阈值库、赢家／合法拒绝与随机前端应作为单独候选接入后再比较；不能把“未启用”写成Ising硬件为零功耗。',bold=True,size=9.5)

page('五 时间与存储究竟限制在哪里')
h('可以相加的端到端阶段')
table(['阶段','ms/token','总时间占比'],[[x['phase'],f"{x['elapsed_ms']:.6f}",f"{x['elapsed_share_pct']:.3f}%"] for x in d['sequential_phase_groups']],[4.3,1.5,1.4],fs=9)
h('不可相加的资源忙时')
table(['资源','忙时ms','相对TPOT'],[[x['service'],f"{x['service_busy_ms']:.6f}",f"{x['over_TPOT_pct']:.2f}%"] for x in d['overlapping_service_busy_times'][:3]], [4.3,1.5,1.4],fs=9)
p('普通MAC忙时远小于投影阶段总时间，说明供数和等待很重要。DMA、SRAM和计算之间存在重叠，上表不是串行时间分解，也不能据此把DMA忙时直接当作不可突破的整机下限。',size=9.5)
h('片上SRAM分布')
table(['物理存储','容量KiB'],[[x['memory'],f"{x['capacity_KiB']:.1f}"] for x in d['memory_capacity']],[5.4,1.8],fs=9)
p(f"共享SRAM显式地址空间保留峰值{d['shared_SRAM_explicit_reserved_peak_bytes']:,}B，约0.682MiB；它不包含独立本地和响应缓冲。8MiB容量与bank／带宽绑定，减少容量不是免费维持原带宽。",size=9.5)
h('外存原粗预算的主要来源')
keep=['weights_physical_bytes','K_selected_physical_bytes','V_sparse_union_physical_bytes','summary_read_physical_bytes']
table(['外存物理读取','MiB/token','64pJ/B代理mJ'],[[x['kind'].replace('_physical_bytes',''),f"{x['MiB']:.3f}",f"{x['legacy_proxy_mJ']:.4f}"] for x in d['external_traffic_proxy'] if x['kind'] in keep],[3.4,1.8,2.0],fs=9)

page('六 证据强度与下一步优化顺序')
h('三个已有门控情景')
table(['原预算情景','片上W','片上mJ/token'],[[label,f"{d['original_scenarios'][key]['covered_onchip_W']:.4f}",f"{d['original_scenarios'][key]['covered_onchip_J']*1e3:.3f}"] for key,label in [('new_counter_CE','新计数器CE 数字CE'),('counter_real_ICG','计数器真实ICG 数字CE'),('counter_ICG_plus_digital_gate_candidate','计数器ICG 数字90%空闲动态消除情景')]],[4.4,1.4,1.4],fs=9)
p('2.516W不是已实测整片：数字门控尚为情景；500MHz按固定电压、每参考拍动态能量向外推，漏电按时间计。原时序没有重新加入这一数字门控情景的唤醒延迟。')
h('优化优先级')
for txt in [
 '能量：先看共享MAC背景、FFN计数与恢复、外存单价。应先核对真实门控和HBM，再判断局部电路改进对整机的收益。',
 '面积：优先计数器与共享SRAM；保持位宽、并行度、吞吐和质量口径，不能把OR4或双随机小核直接替代原多档精确核。',
 '时间：先分供数等待与计算忙时。QKV／O阶段耗时约16.36ms，但MAC算术只占其中一小部分；增加MAC数量不一定改善等待。',
 '缺项：补定点编码、宽定点／FP32到BF16、CIM并行外围和控制／CTS。随机源与比较器拆价也必须匹配已集成边界。',
]:p(txt)
h('已有结果和仍用代理的部分')
p('已有：100MHz计数／X8门控、单数字MAC与FP32加、随机端点、8位LUT及SRAM宏表。代理：B4→B7位数比例、旧恢复、宽累加、数字阵列空闲门控与500MHz实现。未闭合：CIM全并行接口／互连／CTS、部分编码转换、有效HBM能量。')
h('版本与复算')
p('基线main：'+d['base_commit'],size=8.8)
p('原PR192科学锚点：'+d['source_tlm_branch_head'],size=8.8)
p('六个completed摘要原字节SHA、manifest及配置引用已核对；原raw结果SHA只核对封存引用一致，未重新下载大轨迹或门级原件。四配置×三情景复算一致，14模块与阶段时间均闭合。无新PPL、EDA或整机时序执行。',size=9.4)
link('原预算与电路映射说明','https://github.com/rayggxia-workspace/pbit-llm-qwen-baselines/blob/'+d['source_tlm_branch_head']+'/architecture_tlm/ppa_recalibration_20261027/RESULTS_ZH.md')
link('原固定资源调度结果','https://github.com/rayggxia-workspace/pbit-llm-qwen-baselines/blob/'+d['source_tlm_branch_head']+'/architecture_tlm/scheduling_v2_20261027/FINAL_RESULTS_ZH.md')
link('Issue210与本次交付','https://github.com/rayggxia-workspace/pbit-llm-qwen-baselines/issues/210')
p('配套CSV／JSON保留所有逐模块精确值、阶段时间、资源忙时、功能MAC增量、外存事务和未计项；不是把现有完整TLM替换成另一套模型。',size=9.2)
for root in [doc.styles.element,doc.element]:
 for x in list(root.xpath('//w:pBdr')):x.getparent().remove(x)
doc.core_properties.title='整机TLM功耗面积与时间分项'
doc.core_properties.subject='Issue210 原整机预算占比审计'
doc.core_properties.author='pbit LLM研究资料整理'
args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output)
print(str(args.output));print('tables',len(doc.tables),'hyperlinks',len(doc.element.xpath('//w:hyperlink')))
